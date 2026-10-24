import io
from decimal import Decimal
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook
from .models import (
    MateriaPrima,
    OrdenProduccion,
    MovimientoAlmacenMP,
    ManoObraDirecta,
    GastoIndirecto,
    ProrrateoGI,
    ArticuloTerminado,
)


class ContabilidadFlujoCompletoTests(TestCase):
    def test_inicia_vacio_por_defecto(self):
        self.assertEqual(MateriaPrima.objects.count(), 0)
        self.assertEqual(OrdenProduccion.objects.count(), 0)
        self.assertEqual(MovimientoAlmacenMP.objects.count(), 0)
        self.assertEqual(GastoIndirecto.objects.count(), 0)

    def test_normalizacion_kardex_y_borrado_de_captura_o_tarjeta(self):
        mp = MateriaPrima.objects.create(
            codigo="  mp 001  ",
            nombre="  madera de pino  ",
            unidad_medida="PZA",
            email_proveedor="  VENTAS@PROVEEDOR.COM ",
        )
        self.assertEqual(mp.codigo, "MP-001")
        self.assertEqual(mp.nombre, "Madera de pino")
        self.assertEqual(mp.email_proveedor, "ventas@proveedor.com")

        op = OrdenProduccion.objects.create(
            codigo=" op 1050 ",
            producto_terminado=" escritorio ejecutivo ",
            cantidad_unidades=10,
        )
        self.assertEqual(op.codigo, "OP-1050")

        mov1 = MovimientoAlmacenMP.objects.create(
            materia_prima=mp,
            tipo_movimiento="SALDO_INICIAL",
            referencia=" inv 01 ",
            unidades_entrada=Decimal("100.00"),
            costo_unitario=Decimal("50.00"),
        )
        mov2 = MovimientoAlmacenMP.objects.create(
            materia_prima=mp,
            tipo_movimiento="COMPRA",
            referencia=" fac 02 ",
            unidades_entrada=Decimal("100.00"),
            costo_unitario=Decimal("70.00"),
        )
        self.assertEqual(mp.existencia_actual, Decimal("200.00"))
        self.assertEqual(mp.saldo_actual, Decimal("12000.00"))

        # Borrar una captura (mov2) y verificar que la tarjeta recalcula su saldo automáticamente
        resp_del_mov = self.client.post(reverse("eliminar_movimiento_mp", args=[mov2.id]))
        self.assertEqual(resp_del_mov.status_code, 302)
        self.assertEqual(mp.existencia_actual, Decimal("100.00"))
        self.assertEqual(mp.saldo_actual, Decimal("5000.00"))

        # Borrar la tarjeta completa
        resp_del_mp = self.client.post(reverse("eliminar_materia_prima", args=[mp.id]))
        self.assertEqual(resp_del_mp.status_code, 302)
        self.assertEqual(MateriaPrima.objects.count(), 0)

    def test_cargar_preguardados_exportar_excel_con_graficas_y_limpiar(self):
        # 1. Cargar datos preguardados
        resp = self.client.post(reverse("cargar_ejemplo_pizarron"))
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(OrdenProduccion.objects.count(), 3)

        # 2. Verificar Apartado de Resultados y modo oscuro/claro
        r_res = self.client.get(reverse("resultados"))
        self.assertEqual(r_res.status_code, 200)
        self.assertContains(r_res, "canvasPastelElementos")
        self.assertContains(r_res, "btnToggleTheme")
        self.assertContains(r_res, reverse("exportar_resultados_excel"))

        # 3. Exportar a Excel (.xlsx) y verificar que contiene hojas y gráficas nativas
        r_xlsx = self.client.get(reverse("exportar_resultados_excel"))
        self.assertEqual(r_xlsx.status_code, 200)
        self.assertEqual(
            r_xlsx["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb = load_workbook(io.BytesIO(r_xlsx.content))
        self.assertIn("Resultados y Graficas", wb.sheetnames)
        self.assertIn("Tarjetas de Almacen", wb.sheetnames)
        ws = wb["Resultados y Graficas"]
        self.assertGreaterEqual(len(ws._charts), 3)

        # 4. Limpiar datos para volver al estado vacío por defecto
        resp_clean = self.client.post(reverse("limpiar_datos"))
        self.assertEqual(resp_clean.status_code, 302)
        self.assertEqual(OrdenProduccion.objects.count(), 0)
        self.assertEqual(MateriaPrima.objects.count(), 0)

    def test_guardar_mano_de_obra_en_orden_y_sin_duplicado_gi_en_captura(self):
        # Verificar que el apartado de captura no contiene el duplicado de G.I.
        r_cap = self.client.get(reverse("captura"))
        self.assertEqual(r_cap.status_code, 200)
        self.assertNotContains(r_cap, 'value="gi_captura"')

        # Crear una Orden de Producción
        op = OrdenProduccion.objects.create(
            codigo="OP-1050",
            producto_terminado="Mesa Industrial",
            cantidad_unidades=10,
        )

        # 1. Guardar Mano de Obra desde Captura con importe directo (dejando cuota_por_hora vacío)
        resp_mo_cap = self.client.post(
            reverse("captura"),
            {
                "tipo_captura": "mo_captura",
                "mo-orden_produccion": op.id,
                "mo-fecha": "2026-10-23",
                "mo-referencia_nomina": "NOM-01",
                "mo-operario_departamento": "Ensamble",
                "mo-horas_trabajadas": "40",
                "mo-cuota_por_hora": "",
                "mo-importe_total": "5000.00",
            },
        )
        self.assertEqual(resp_mo_cap.status_code, 302)
        self.assertEqual(op.cargos_mo.count(), 1)
        self.assertEqual(op.total_mo, Decimal("5000.00"))

        # 2. Guardar Mano de Obra directamente dentro de Órdenes de Producción (por horas x cuota)
        resp_mo_op = self.client.post(
            reverse("ordenes_produccion"),
            {
                "accion": "agregar_mo_orden",
                "mo_op-orden_produccion": op.id,
                "mo_op-fecha": "2026-10-23",
                "mo_op-referencia_nomina": "NOM-02",
                "mo_op-operario_departamento": "Acabado",
                "mo_op-horas_trabajadas": "20",
                "mo_op-cuota_por_hora": "150.00",
                "mo_op-importe_total": "",
            },
        )
        self.assertEqual(resp_mo_op.status_code, 302)
        self.assertEqual(op.cargos_mo.count(), 2)
        self.assertEqual(op.total_mo, Decimal("8000.00"))

        # Verificar que se muestra dentro del detalle de la orden en /ordenes-produccion/
        r_op = self.client.get(f"{reverse('ordenes_produccion')}?op_id={op.id}")
        self.assertEqual(r_op.status_code, 200)
        self.assertContains(r_op, "NOM-01")
        self.assertContains(r_op, "NOM-02")
        self.assertContains(r_op, "$8000.00")
